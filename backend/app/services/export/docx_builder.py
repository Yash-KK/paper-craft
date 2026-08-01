"""Build styled question-paper DOCX files from generation JSON.

The built-in sample DOCX (``samples/40_marks_sample.docx``) is used only for
page layout / styles. All visible content comes from Paper Details and the
generated paper JSON. Requires ``pandoc`` on PATH for LaTeX → Word equations.
"""

from __future__ import annotations

import copy
import io
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any

import docx
from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.run import Run

from app.services.export.latex import normalize_newlines, prepare_markdown_for_pandoc
from app.services.export.section_copy import (
    format_option_label,
    format_section_heading,
    infer_section_question_type,
    options_should_be_single_line,
    question_body_lines,
    resolve_final_paper,
    section_description,
    section_marks_summary,
    strip_embedded_options,
)

FONT_NAME = "Bookman Old Style"
SIZE_SCHOOL_NAME = 16
SIZE_EXAM_TITLE = 14
SIZE_BODY = 12
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
INLINE_OPTION_TAB_STOPS_CM = (3.8, 7.2, 10.6)
QUESTION_BODY_INDENT_CM = 0.5
QUESTION_CONTINUATION_INDENT_CM = 0.75

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class PandocNotFoundError(RuntimeError):
    """Raised when the pandoc binary is missing from PATH."""


def ensure_pandoc() -> None:
    if shutil.which("pandoc") is None:
        raise PandocNotFoundError(
            "pandoc is required to export DOCX with equations. "
            "Install pandoc and ensure it is on PATH."
        )


def _pandoc_markdown_to_document_xml(md_text: str) -> bytes:
    ensure_pandoc()
    try:
        result = subprocess.run(
            ["pandoc", "-f", "markdown", "-t", "docx", "-o", "-"],
            input=md_text.encode("utf-8"),
            capture_output=True,
            check=True,
        )
    except FileNotFoundError as exc:
        raise PandocNotFoundError(
            "pandoc is required to export DOCX with equations."
        ) from exc
    except subprocess.CalledProcessError as exc:
        detail = (exc.stderr or b"").decode("utf-8", errors="replace").strip()
        raise RuntimeError(f"pandoc failed: {detail or exc}") from exc

    with zipfile.ZipFile(io.BytesIO(result.stdout)) as archive:
        return archive.read("word/document.xml")


def _etree_localname(element) -> str:
    from lxml import etree

    return etree.QName(element).localname


def _extract_first_paragraph_children(document_xml: bytes):
    root = parse_xml(document_xml)
    body = root.find(f"{{{W_NS}}}body")
    paragraph = body.find(f"{{{W_NS}}}p")
    if paragraph is None:
        return []
    return [child for child in paragraph if _etree_localname(child) != "pPr"]


def _tighten_spacing(paragraph):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return paragraph


def add_blank_line(doc):
    return _tighten_spacing(doc.add_paragraph())


def add_rich_paragraph(
    doc,
    text,
    bold=False,
    italic=False,
    size_pt=SIZE_BODY,
    alignment=None,
    font_name=FONT_NAME,
    indent_cm=None,
):
    paragraph = doc.add_paragraph()
    _tighten_spacing(paragraph)
    if alignment is not None:
        paragraph.alignment = alignment
    if indent_cm is not None:
        paragraph.paragraph_format.left_indent = docx.shared.Cm(indent_cm)

    if not text.strip():
        return paragraph

    md = prepare_markdown_for_pandoc(text)
    children = _extract_first_paragraph_children(_pandoc_markdown_to_document_xml(md))
    for child in children:
        from lxml import etree

        tag = etree.QName(child).localname
        el = copy.deepcopy(child)
        if tag == "r":
            run_obj = Run(el, paragraph)
            run_obj.font.name = font_name
            run_obj.font.size = Pt(size_pt)
            run_obj.bold = bold
            run_obj.italic = italic
        paragraph._p.append(el)
    return paragraph


def add_rich_block(doc, text, **kwargs):
    if not text:
        return
    text = normalize_newlines(text)
    for line in re.split(r"\n+", text):
        line = line.strip()
        if line:
            add_rich_paragraph(doc, line, **kwargs)


def load_template(template_docx: str | Path) -> Document:
    """Keep styles / page setup from the template; strip all body content."""
    doc = docx.Document(str(template_docx))
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)
    return doc


def add_bottom_border(paragraph, size=12, color="000000"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = parse_xml(
        f'<w:pBdr xmlns:w="{W_NS}">'
        f'<w:bottom w:val="single" w:sz="{size}" w:space="1" w:color="{color}"/>'
        f"</w:pBdr>"
    )
    p_pr.append(p_bdr)


def set_list_numbering(doc, paragraph, num_id: int, ilvl: int = 0):
    paragraph.style = doc.styles["List Paragraph"]
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = parse_xml(
        f'<w:numPr xmlns:w="{W_NS}">'
        f'<w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/>'
        f"</w:numPr>"
    )
    p_pr.append(num_pr)


def _get_paper_field(question_paper: Any, key: str, default: Any = None) -> Any:
    if isinstance(question_paper, dict):
        return question_paper.get(key, default)
    return getattr(question_paper, key, default)


def header_from_question_paper(question_paper: Any) -> dict[str, Any]:
    exam_date = _get_paper_field(question_paper, "exam_date")
    if exam_date is not None and hasattr(exam_date, "isoformat"):
        exam_date = exam_date.isoformat()
    return {
        "school_name": _get_paper_field(question_paper, "school_name"),
        "exam_title": _get_paper_field(question_paper, "exam_title"),
        "subject": _get_paper_field(question_paper, "subject"),
        "grade": _get_paper_field(question_paper, "grade"),
        "total_marks": _get_paper_field(question_paper, "total_marks"),
        "duration_minutes": _get_paper_field(question_paper, "duration_minutes"),
        "exam_date": exam_date,
        "general_instructions": _get_paper_field(question_paper, "general_instructions")
        or [],
    }


def _format_duration(duration_minutes: Any) -> str:
    if not duration_minutes:
        return ""
    hours = float(duration_minutes) / 60
    return f"DURATION: {hours:g} HOUR(S)"


def add_header_block(
    doc,
    header: dict,
    title_suffix: str | None = None,
):
    """Write paper-details header only — never copy text from the sample DOCX."""
    school_name = (header.get("school_name") or "").strip()
    if school_name:
        add_rich_paragraph(
            doc,
            school_name,
            bold=True,
            size_pt=SIZE_SCHOOL_NAME,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    title = (header.get("exam_title") or "").strip()
    if title_suffix:
        title = f"{title} - {title_suffix}" if title else title_suffix
    if title:
        add_rich_paragraph(
            doc,
            title,
            bold=True,
            size_pt=SIZE_EXAM_TITLE,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    grade = header.get("grade")
    grade_txt = "" if grade is None else str(grade)
    add_rich_paragraph(
        doc,
        f"NAME: {'…' * 18}   ROLL NO: {'…' * 6}   GRADE: {grade_txt}",
        bold=True,
        size_pt=SIZE_BODY,
    )

    subject = (header.get("subject") or "").strip()
    exam_date = (header.get("exam_date") or "").strip()
    subject_bits = [f"SUBJECT: {subject}" if subject else "SUBJECT:"]
    if exam_date:
        subject_bits.append(f"DATE: {exam_date}")
    subject_line = add_rich_paragraph(
        doc,
        "          ".join(subject_bits),
        bold=True,
        size_pt=SIZE_BODY,
    )
    add_bottom_border(subject_line)

    duration_txt = _format_duration(header.get("duration_minutes"))
    total_marks = header.get("total_marks")
    marks_txt = "" if total_marks is None else str(total_marks)
    marks_bits = [bit for bit in [duration_txt, f"MAX MARKS: {marks_txt}"] if bit]
    marks_line = add_rich_paragraph(
        doc,
        "          ".join(marks_bits),
        bold=True,
        size_pt=SIZE_BODY,
    )
    add_bottom_border(marks_line)

    instructions = [
        str(instr).strip()
        for instr in (header.get("general_instructions") or [])
        if str(instr).strip()
    ]
    if instructions:
        add_rich_paragraph(doc, "General Instructions:", bold=True, size_pt=SIZE_BODY)
        add_rich_paragraph(
            doc,
            "Read the following instructions carefully and strictly follow them:",
            bold=True,
            size_pt=SIZE_BODY,
        )
        for instr in instructions:
            paragraph = add_rich_paragraph(doc, instr, size_pt=SIZE_BODY)
            set_list_numbering(doc, paragraph, num_id=1)
        add_blank_line(doc)


def add_section_header(
    doc,
    section_name: str,
    section_questions: list[dict],
    section_instructions: str | None = None,
    question_type: str | None = None,
):
    add_blank_line(doc)
    add_rich_paragraph(
        doc,
        format_section_heading(section_name),
        bold=True,
        size_pt=SIZE_BODY,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_blank_line(doc)

    count, marks_each, total = section_marks_summary(section_questions)
    if section_instructions and section_instructions.strip():
        description = section_instructions.strip()
        if marks_each is not None and "×" not in description:
            description = f"{description}      {count} × {marks_each:g} = {total:g}M"
    else:
        description = section_description(
            question_type=question_type
            or infer_section_question_type(section_questions),
            question_count=count,
            marks_each=marks_each,
            total_marks=total,
        )
    add_rich_paragraph(
        doc,
        description,
        bold=True,
        size_pt=SIZE_BODY,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )


def add_plain_paragraph(
    doc,
    text,
    *,
    size_pt=SIZE_BODY,
    font_name=FONT_NAME,
    indent_cm=None,
):
    paragraph = doc.add_paragraph()
    _tighten_spacing(paragraph)
    if indent_cm is not None:
        paragraph.paragraph_format.left_indent = Cm(indent_cm)
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    return paragraph


def add_inline_options_paragraph(
    doc, options: list[str], *, indent_cm=QUESTION_BODY_INDENT_CM
):
    """Render MCQ options on one line with tab stops (pandoc collapses spaces)."""
    paragraph = doc.add_paragraph()
    _tighten_spacing(paragraph)
    paragraph.paragraph_format.left_indent = Cm(indent_cm)
    tab_stops = paragraph.paragraph_format.tab_stops
    for position in INLINE_OPTION_TAB_STOPS_CM:
        tab_stops.add_tab_stop(Cm(position))

    for index, option in enumerate(options):
        if index > 0:
            paragraph.add_run("\t")
        run = paragraph.add_run(format_option_label(option, index))
        run.font.name = FONT_NAME
        run.font.size = Pt(SIZE_BODY)
    return paragraph


def add_question(doc, q: dict, case_study_number: int | None = None):
    if case_study_number is not None:
        add_rich_paragraph(
            doc,
            f"Case Study - {case_study_number}",
            bold=True,
            size_pt=SIZE_BODY,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )

    question_text = q.get("question_text") or ""
    options = q.get("options") or []
    if options:
        question_text = strip_embedded_options(question_text)

    body_lines = question_body_lines(q.get("question_number"), question_text)
    for index, line in enumerate(body_lines):
        indent_cm = (
            QUESTION_BODY_INDENT_CM if index == 0 else QUESTION_CONTINUATION_INDENT_CM
        )
        add_rich_paragraph(doc, line, size_pt=SIZE_BODY, indent_cm=indent_cm)

    if options:
        if options_should_be_single_line(q.get("question_type")):
            add_inline_options_paragraph(doc, options)
        else:
            for i, opt in enumerate(options):
                add_plain_paragraph(
                    doc,
                    format_option_label(opt, i),
                    indent_cm=QUESTION_BODY_INDENT_CM,
                )

    if q.get("alternate_question_text"):
        add_rich_paragraph(
            doc,
            "(OR)",
            bold=True,
            size_pt=SIZE_BODY,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
        alt = q["alternate_question_text"]
        if options:
            alt = strip_embedded_options(alt)
        add_rich_block(doc, alt, size_pt=SIZE_BODY)

    add_blank_line(doc)


def section_instructions_map(question_paper: Any) -> dict[str, str | None]:
    sections = _get_paper_field(question_paper, "sections") or []
    result: dict[str, str | None] = {}
    for section in sections:
        if isinstance(section, dict):
            name, instr = (
                section.get("section_name"),
                section.get("section_instructions"),
            )
        else:
            name = getattr(section, "section_name", None)
            instr = getattr(section, "section_instructions", None)
        if name:
            result[name] = instr
    return result


def section_question_type_map(question_paper: Any) -> dict[str, str | None]:
    sections = _get_paper_field(question_paper, "sections") or []
    result: dict[str, str | None] = {}
    for section in sections:
        if isinstance(section, dict):
            name = section.get("section_name")
            qtype = section.get("question_type")
        else:
            name = getattr(section, "section_name", None)
            qtype = getattr(section, "question_type", None)
        if name:
            if qtype is not None and hasattr(qtype, "value"):
                qtype = qtype.value
            result[str(name)] = str(qtype) if qtype else None
    return result


def ordered_section_names(question_paper: Any, assembled_sections: dict) -> list[str]:
    """Prefer Paper Details / marking-scheme order; append any extras last."""
    blueprint_sections = _get_paper_field(question_paper, "sections") or []
    ordered: list[str] = []
    seen: set[str] = set()
    for section in blueprint_sections:
        name = (
            section.get("section_name")
            if isinstance(section, dict)
            else getattr(section, "section_name", None)
        )
        if name and name not in seen:
            ordered.append(name)
            seen.add(name)
    for name in assembled_sections:
        if name not in seen:
            ordered.append(name)
            seen.add(name)
    return ordered


def _build_question_paper_document(
    question_paper: Any,
    final_paper: dict,
    template_docx: str | Path,
) -> Document:
    doc = load_template(template_docx)
    add_header_block(doc, header_from_question_paper(question_paper))

    sections = resolve_final_paper(final_paper).get("sections") or {}
    instructions_by_section = section_instructions_map(question_paper)
    type_by_section = section_question_type_map(question_paper)
    for section_name in ordered_section_names(question_paper, sections):
        questions = sections.get(section_name) or []
        if not questions:
            continue
        add_section_header(
            doc,
            section_name,
            questions,
            instructions_by_section.get(section_name),
            question_type=type_by_section.get(section_name),
        )
        case_study_counter = 0
        for question in questions:
            case_study_number = None
            if question.get("question_type") == "CASE_STUDY":
                case_study_counter += 1
                case_study_number = case_study_counter
            add_question(doc, question, case_study_number=case_study_number)
    return doc


def _document_to_bytes(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_question_paper_docx(
    question_paper: Any,
    final_paper: dict,
    template_docx: str | Path,
    out_path: str | Path | None = None,
) -> Path | bytes:
    doc = _build_question_paper_document(question_paper, final_paper, template_docx)
    if out_path is None:
        return _document_to_bytes(doc)
    path = Path(out_path)
    doc.save(str(path))
    return path


def render_question_paper_docx_bytes(
    question_paper: Any,
    final_paper: dict,
    template_docx: str | Path,
) -> bytes:
    return _document_to_bytes(
        _build_question_paper_document(question_paper, final_paper, template_docx)
    )


def export_question_paper_only(
    question_paper: Any,
    final_paper: dict,
    template_docx: str | Path,
    paper_out: str | Path = "generated_question_paper.docx",
) -> Path:
    result = build_question_paper_docx(
        question_paper, final_paper, template_docx, paper_out
    )
    assert isinstance(result, Path)
    return result
